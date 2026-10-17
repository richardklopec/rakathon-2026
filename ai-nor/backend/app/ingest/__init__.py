"""Turn uploaded files into plain-text Documents. Offsets in Evidence refer to Document.text."""

import io
import re
import uuid
from datetime import datetime, timedelta

import openpyxl

from app.models import Document

_KIND_BY_NAME = [
    (r"histolog", "histology"),
    (r"^pz|propoust", "PZ"),
    (r"^az|ambulant", "AZ"),
    (r"^kz|mdt", "MDT"),
    (r"operac", "operation"),
]


def guess_kind(filename: str) -> str:
    name = filename.lower()
    for pattern, kind in _KIND_BY_NAME:
        if re.search(pattern, name):
            return kind
    return "other"


def _normalise(text: str) -> str:
    text = text.replace("\r\n", "\n").replace(" ", " ")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def _excel_date(v):
    # NIS exports store dates as Excel serial numbers
    if isinstance(v, (int, float)) and 30000 < v < 80000:
        return (datetime(1899, 12, 30) + timedelta(days=v)).strftime("%Y-%m-%d %H:%M")
    return v


def xlsx_to_documents(data: bytes, filename: str) -> list[Document]:
    """NIS-style export: one row = one document, first row = column names."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = list(ws.iter_rows(values_only=True))
    header, docs = [str(h) for h in rows[0]], []
    for i, row in enumerate(rows[1:], 1):
        if not any(row):
            continue
        parts = []
        for col, value in zip(header, row):
            if value in (None, ""):
                continue
            value = _excel_date(value) if "date" in col.lower() else value
            value = re.sub(r"<br\s*/?>", "\n", str(value))
            value = re.sub(r"</?body>", "", value)
            parts.append(f"{col}: {value}")
        docs.append(
            Document(
                id=str(uuid.uuid4()),
                filename=f"{filename}#{i}",
                kind=guess_kind(str(row[0] or filename)),
                text=_normalise("\n".join(parts)),
            )
        )
    return docs


def to_documents(filename: str, data: bytes) -> list[Document]:
    from app.ingest import formats  # scans (OCR), Word, e-mail

    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "xlsx":
        return xlsx_to_documents(data, filename)
    if suffix == "eml":
        return formats.eml_to_documents(data, filename)
    if suffix == "pdf":
        text = formats.pdf_or_image_to_text(data)  # text layer, OCR for scanned pages
    elif suffix in formats.IMAGE_TYPES:
        text = formats.pdf_or_image_to_text(data, formats.IMAGE_TYPES[suffix])
    elif suffix == "docx":
        text = formats.docx_to_text(data)
    elif suffix in ("txt", "md"):
        text = data.decode("utf-8")
    else:
        raise ValueError(f"Nepodporovaný typ souboru: {filename} (podporováno: {', '.join(formats.SUPPORTED)})")
    return [Document(id=str(uuid.uuid4()), filename=filename, kind=guess_kind(filename), text=_normalise(text))]
